from django.core.mail import send_mail
from django.conf import settings


def get_device_details(headers, token):
    device_name = headers.get("HTTP_X_DEVICE_MODEL")
    user_agent = headers.get("HTTP_USER_AGENT", "")
    if not device_name:
        device_name = user_agent
        device_details = ""
    else:
        device_details = user_agent

    return device_name, device_details


def get_ip_address(request):
    x_forwarded_for = request.META.get("HTTP_X_FORWARDED_FOR")
    if x_forwarded_for:
        ip = x_forwarded_for.split(",")[0]
    else:
        ip = request.META.get("REMOTE_ADDR")
    return ip


def warning_mail_send(user_email, ip_address):
    text_content = (
        f"we recorded an attempt to log in to your account with this ip address {ip_address}. If it was you, then simply"
        "ignore this letter. If it wasn’t you - contact the site administration as soon as possible."
    )
    subject = "Attempt to log in to your account"
    from_email = settings.DEFAULT_FROM_EMAIL
    # receive warning mail
    mail_sent = send_mail(subject, text_content, from_email, [user_email])

    return mail_sent


import csv
import json
import openpyxl
import docx
from PyPDF2 import PdfFileReader
import io


def parse_word_document(stream):
    """
    Parse a Word document and return the data.
    """
    # Use python-docx to parse the Word document
    document = docx.Document(stream)
    data = []
    for paragraph in document.paragraphs:
        data.append(paragraph.text)
    return str(data)


def parse_json(stream):
    """
    Parse a JSON document and return the data.
    """
    # Use the json module to parse the JSON document
    data = json.loads(stream.read())
    return str(data)


def parse_pdf(stream):
    """
    Parse a PDF and return the data.
    """
    # Use PyPDF2 to parse the PDF
    pdf = PdfFileReader(stream)
    data = []
    for page in pdf.pages:
        data.append(page.extractText())
    return str(data)


def parse_csv(stream):
    """
    Parse a CSV file and return the data.
    """
    # Use the csv module to parse the CSV file
    data = []
    # Wrap the stream in a TextIOWrapper to read it as strings
    reader = csv.reader(io.TextIOWrapper(stream, encoding="utf-8"))
    for row in reader:
        data.append(row)
    return str(data)


def parse_excel(stream):
    """
    Parse an Excel spreadsheet and return the data.
    """
    # Use openpyxl to parse the Excel spreadsheet
    workbook = openpyxl.load_workbook(stream)
    data = []
    for sheet in workbook:
        for row in sheet.rows:
            data.append([cell.value for cell in row])
    return str(data)
